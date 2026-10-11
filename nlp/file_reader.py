"""
file_reader.py  (nlp folder)
Candidate CVs (PDF / Word / TXT) se text nikalta hai.
Is text ko aage extractor.py ko diya jata hai (skills, experience, education).

Install:  pip install PyPDF2 python-docx
"""
import io
import re
from pathlib import Path

from PyPDF2 import PdfReader
from docx import Document

SUPPORTED = {".pdf", ".docx", ".txt"}


def clean_text(text: str) -> str:
    """Extra spaces/blank lines hatao, line structure (sections) bacha kar."""
    text = text.replace("\x00", " ").replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.splitlines()).strip()


def read_pdf(source) -> str:
    """source: file path ya file-like object (BytesIO)."""
    reader = PdfReader(source)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("PDF password-protected hai")
    pages = [(page.extract_text() or "") for page in reader.pages]
    return clean_text("\n".join(pages))


def read_docx(source) -> str:
    """Paragraphs ke saath tables ka text bhi nikalta hai."""
    doc = Document(source)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return clean_text("\n".join(parts))


def read_txt(source) -> str:
    data = Path(source).read_bytes() if not hasattr(source, "read") else source.read()
    return clean_text(data.decode("utf-8", errors="ignore"))


def read_cv(path) -> str:
    """File path se text. Extension dekh kar sahi reader chunta hai."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {ext} (sirf PDF, DOCX, TXT)")
    if ext == ".pdf":
        text = read_pdf(str(path))
    elif ext == ".docx":
        text = read_docx(str(path))
    else:
        text = read_txt(path)
    if len(text) < 30:
        raise ValueError(f"{path.name}: text nahi mila (shayad scanned/image PDF hai, OCR chahiye)")
    return text


def read_cv_bytes(content: bytes, filename: str) -> str:
    """FastAPI upload ke liye: await file.read() ke bytes + filename do."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {ext}")
    buf = io.BytesIO(content)
    text = read_pdf(buf) if ext == ".pdf" else read_docx(buf) if ext == ".docx" else clean_text(content.decode("utf-8", "ignore"))
    if len(text) < 30:
        raise ValueError(f"{filename}: text nahi mila (scanned PDF ho sakta hai)")
    return text


def read_folder(folder) -> dict:
    """Folder ki saari CVs padho. Return: {filename: text}. Kharab files skip + report."""
    results, failed = {}, {}
    for p in sorted(Path(folder).iterdir()):
        if p.suffix.lower() in SUPPORTED:
            try:
                results[p.name] = read_cv(p)
            except Exception as e:
                failed[p.name] = str(e)
    for name, err in failed.items():
        print(f"[SKIPPED] {name}: {err}")
    return results


if __name__ == "__main__":
    import sys
    folder = sys.argv[1] if len(sys.argv) > 1 else "../data/cvs"
    cvs = read_folder(folder)
    print(f"{len(cvs)} CVs read hui")
    for name, text in cvs.items():
        print(f"- {name}: {len(text)} characters")
